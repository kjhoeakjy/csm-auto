"""PPT 그리기 도구 (텍스트, 선, 도형, 선 아이콘)와 글꼴 규칙
- 흰 바탕, 검정 글자, 회색 2단계. 박스·그림자·그라데이션 없음, 1px 선과 여백으로 구분
- 글꼴: Pretendard 정적 글꼴 (Regular / Medium / SemiBold / Bold), pptx·PDF 에 포함
"""
import math
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

HERE = Path(__file__).resolve().parent
OUT = HERE.parent / "역량PT_김연주.pptx"

INK = RGBColor(0x00, 0x00, 0x00)
G1 = RGBColor(0x6B, 0x6B, 0x6B)      # 회색 1: 키워드·라벨
G2 = RGBColor(0x9A, 0x9A, 0x9A)      # 회색 2: 연월·각주
LINE = RGBColor(0xDD, 0xDD, 0xDD)    # 구분선
BAND = RGBColor(0xF6, 0xF6, 0xF6)    # 단계 띠 바탕
ACCENT = RGBColor(0xC8, 0x10, 0x2E)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
# 정적 글꼴(Pretendard / Medium / SemiBold, Bold 는 굵게 속성). PDF 에 글꼴이 들어가도록 가변 글꼴 대신 사용.
# 이 PC 에는 가변 글꼴만 설치돼 있어 export.ps1 이 정적 인스턴스를 만들어 이번 세션에만 불러옴 (fonts.py)
F = {"reg": "Pretendard", "med": "Pretendard Medium", "sb": "Pretendard SemiBold", "bold": "Pretendard"}
M = 0.6                               # 좌우 여백 (inch)
W, H = 13.333, 7.5
PX = 0.75                             # 1px 선 (pt)
ICON = 1.5                            # 아이콘·트랙 선 굵기 (pt)


# ---------------------------------------------------------------- 그리기 도구
def text(slide, x, y, w, h, runs, size=11, color=INK, weight="reg", align=PP_ALIGN.LEFT,
         anchor=MSO_ANCHOR.TOP, spacing=1.15):
    """runs: 문자열 하나 또는 줄 목록(줄 = [(문자열, {size, color, weight}), ...]).
    줄바꿈은 직접 지정 (자동 줄바꿈 끔 → 단어 중간 줄바꿈 없음)"""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    lines = runs if isinstance(runs, list) else [[(runs, {})]]
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.line_spacing = spacing
        for t, o in (line if isinstance(line, list) else [(line, {})]):
            r = p.add_run()
            r.text = t
            wt = o.get("weight", weight)
            r.font.size = Pt(o.get("size", size))
            r.font.color.rgb = o.get("color", color)
            r.font.name = F[wt]
            r.font.bold = wt == "bold"
            rpr = r._r.get_or_add_rPr()          # 한글 글꼴도 같은 글꼴로
            for tag in ("a:ea", "a:cs"):
                el = rpr.find(qn(tag))
                if el is None:
                    el = rpr.makeelement(qn(tag), {})
                    rpr.append(el)
                el.set("typeface", F[wt])
    return tb


def _plain(shape):
    shape.shadow.inherit = False
    return shape


def rect(slide, x, y, w, h, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    return _plain(r)


def hline(slide, x, y, w, color=LINE, width=PX):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y) - Pt(width) / 2, Inches(w), Pt(width))
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    return _plain(r)


def vline(slide, x, y, h, color=LINE, width=PX):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x) - Pt(width) / 2, Inches(y), Pt(width), Inches(h))
    r.fill.solid(); r.fill.fore_color.rgb = color
    r.line.fill.background()
    return _plain(r)


def seg(slide, x1, y1, x2, y2, color=INK, width=ICON):
    """직선 (테마 스타일 제거 → 그림자 없음)"""
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    st = c._element.find(qn("p:style"))
    if st is not None:
        c._element.remove(st)
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def poly(slide, pts, closed=False, color=INK, width=ICON):
    """꺾은선·다각형 (inch 좌표 목록)"""
    e = [(Inches(x), Inches(y)) for x, y in pts]
    b = slide.shapes.build_freeform(e[0][0], e[0][1], scale=1.0)
    b.add_line_segments(e[1:], close=closed)
    s = b.convert_to_shape()
    s.fill.background()
    s.line.color.rgb = color
    s.line.width = Pt(width)
    return _plain(s)


def circle(slide, cx, cy, d, fill=None, line=INK, width=ICON):
    s = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(width)
    return _plain(s)


def header(slide, tag, title, sub=None):
    text(slide, M, 0.42, 4, 0.25, tag, size=10, color=G1, weight="med")
    text(slide, M, 0.66, W - 2 * M, 0.45, title, size=20, weight="bold")
    if sub:
        text(slide, M, 1.12, W - 2 * M, 0.3, sub, size=11, color=G1)


# ---------------------------------------------------------------- 선 아이콘 (중심 x, y 기준, 반경 약 0.12 inch)
def ic_bar(s, x, y):     # 막대 그래프
    for dx, top in [(-0.075, 0.03), (0, -0.05), (0.075, -0.1)]:
        seg(s, x + dx, y + 0.1, x + dx, y + top)


def ic_flag(s, x, y):    # 깃발
    seg(s, x - 0.08, y + 0.12, x - 0.08, y - 0.12)
    poly(s, [(x - 0.08, y - 0.11), (x + 0.1, y - 0.055), (x - 0.08, y + 0.0)], closed=True)


def ic_net(s, x, y):     # 신경망
    a = [(x - 0.095, y - 0.06), (x - 0.095, y + 0.06)]
    b = [(x, y - 0.1), (x, y), (x, y + 0.1)]
    c = [(x + 0.095, y)]
    for p in a:
        for q in b:
            seg(s, *p, *q, width=0.75)
    for p in b:
        for q in c:
            seg(s, *p, *q, width=0.75)
    for px, py in a + b + c:
        circle(s, px, py, 0.045, fill=WHITE, width=1.25)


def ic_building(s, x, y):  # 건물
    poly(s, [(x - 0.09, y + 0.12), (x - 0.09, y - 0.11), (x + 0.09, y - 0.11), (x + 0.09, y + 0.12)], closed=True)
    for yy in (-0.05, 0.01):
        seg(s, x - 0.045, y + yy, x + 0.045, y + yy, width=1.0)
    seg(s, x, y + 0.06, x, y + 0.12, width=1.0)


def ic_curve(s, x, y):   # 생존곡선 (계단형)
    poly(s, [(x - 0.11, y - 0.12), (x - 0.11, y + 0.11), (x + 0.12, y + 0.11)])
    poly(s, [(x - 0.11, y - 0.09), (x - 0.05, y - 0.09), (x - 0.05, y - 0.03), (x + 0.01, y - 0.03),
             (x + 0.01, y + 0.03), (x + 0.1, y + 0.03)])


def ic_dots(s, x, y):    # 연결된 점
    pts = [(x - 0.085, y + 0.07), (x + 0.085, y + 0.07), (x, y - 0.085)]
    for i in range(3):
        seg(s, *pts[i], *pts[(i + 1) % 3], width=1.0)
    for px, py in pts:
        circle(s, px, py, 0.065, fill=WHITE)


def ic_layers(s, x, y):  # 겹친 층
    poly(s, [(x - 0.11, y - 0.05), (x, y - 0.105), (x + 0.11, y - 0.05), (x, y + 0.005)], closed=True)
    for yy in (0.02, 0.085):
        poly(s, [(x - 0.11, y + yy - 0.025), (x, y + yy + 0.03), (x + 0.11, y + yy - 0.025)])


def ic_gear(s, x, y):    # 톱니
    for k in range(8):
        t = k * math.pi / 4
        seg(s, x + 0.07 * math.cos(t), y + 0.07 * math.sin(t), x + 0.12 * math.cos(t), y + 0.12 * math.sin(t))
    circle(s, x, y, 0.15, fill=WHITE)
    circle(s, x, y, 0.05)


# ---------------------------------------------------------------- 대본 → 발표자 노트
